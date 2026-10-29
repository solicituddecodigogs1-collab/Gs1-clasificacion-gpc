"""Helpers para generar presentaciones GS1 Healthcare (16:9, Arial) con python-pptx.

Uso:
    from utils.gs1_healthcare_pptx import *
    prs = create_gs1_deck()
    add_gs1_cover(prs, "Título", "Subtítulo", "Fecha | Autor")
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_gs1_header(s, "Conclusión principal en una oración completa", slide_num=2)
    prs.save("salida.pptx")
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

GS1_BLUE = RGBColor(0, 44, 108)       # #002C6C primario
GS1_HC_TEAL = RGBColor(0, 133, 122)   # #00857A acento salud
GS1_ORANGE = RGBColor(242, 99, 52)    # #F26334 destacado
GS1_SKY = RGBColor(0, 156, 222)       # #009CDE secundario
GS1_TEXT = RGBColor(51, 65, 85)       # #334155 cuerpo
GS1_BG_LIGHT = RGBColor(240, 244, 248)  # #F0F4F8 tarjetas/tablas
GS1_BORDER = RGBColor(203, 213, 225)  # #CBD5E1 líneas/bordes
WHITE = RGBColor(255, 255, 255)

FONT = "Arial"
SLIDE_W, SLIDE_H = 13.333, 7.5
MAX_BULLETS = 5
MAX_WORDS = 10


def _rect(slide, left, top, width, height, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def _textbox(slide, left, top, width, height, text, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _style(r, size, color, bold)
    return tb


def _style(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _check_bullets(bullets):
    if len(bullets) > MAX_BULLETS:
        raise ValueError(f"Máximo {MAX_BULLETS} viñetas por bloque (recibidas {len(bullets)})")
    for b in bullets:
        if len(b.split()) > MAX_WORDS:
            raise ValueError(f"Viñeta con más de {MAX_WORDS} palabras: {b!r}")


def create_gs1_deck():
    """Presentación vacía 16:9 (13.333 x 7.5 in)."""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def add_gs1_cover(prs, title, subtitle, meta_info):
    """Portada: fondo GS1_BLUE, barra vertical naranja a la izquierda, acento inferior teal."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, GS1_BLUE)
    _rect(slide, 0, 0, 0.35, SLIDE_H, GS1_ORANGE)
    _rect(slide, 0.35, SLIDE_H - 0.3, SLIDE_W - 0.35, 0.3, GS1_HC_TEAL)
    _textbox(slide, 1.0, 0.7, 6.0, 0.4, "GS1 HEALTHCARE", 14, GS1_ORANGE, bold=True)
    _textbox(slide, 1.0, 2.0, 11.0, 2.0, title, 40, WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _rect(slide, 1.05, 4.2, 1.5, 0.06, GS1_ORANGE)
    _textbox(slide, 1.0, 4.4, 11.0, 1.0, subtitle, 22, WHITE)
    _textbox(slide, 1.0, 6.2, 11.0, 0.5, meta_info, 14, GS1_BORDER)
    return slide


def add_gs1_header(slide, action_title, tracker="GS1 HEALTHCARE", slide_num=None):
    """Encabezado: tracker naranja, action title azul, línea doble teal + naranja."""
    _textbox(slide, 0.6, 0.3, 6.0, 0.3, tracker, 11, GS1_ORANGE, bold=True)
    if slide_num is not None:
        _textbox(slide, SLIDE_W - 1.6, 0.3, 1.0, 0.3, str(slide_num), 11, GS1_TEXT,
                 align=PP_ALIGN.RIGHT)
    _textbox(slide, 0.6, 0.6, SLIDE_W - 1.2, 0.9, action_title, 26, GS1_BLUE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    _rect(slide, 0.6, 1.6, SLIDE_W - 1.2, 0.04, GS1_HC_TEAL)
    _rect(slide, 0.6, 1.66, SLIDE_W - 1.2, 0.02, GS1_ORANGE)


def add_gs1_kpi_card(slide, left, top, width, height, value, label, sublabel="", highlight=False):
    """Tarjeta KPI con borde superior de color (naranja si highlight, teal por defecto)."""
    accent = GS1_ORANGE if highlight else GS1_HC_TEAL
    _rect(slide, left, top, width, height, GS1_BG_LIGHT, line=GS1_BORDER)
    _rect(slide, left, top, width, 0.08, accent)
    _textbox(slide, left + 0.1, top + 0.15, width - 0.2, height * 0.42, str(value), 32, accent,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _textbox(slide, left + 0.1, top + 0.15 + height * 0.42, width - 0.2, height * 0.2, label,
             13, GS1_BLUE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if sublabel:
        _textbox(slide, left + 0.1, top + 0.15 + height * 0.62, width - 0.2, height * 0.2, sublabel,
                 10, GS1_TEXT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def add_gs1_content_box(slide, left, top, width, height, title, bullets, accent_color=GS1_HC_TEAL):
    """Caja con barra lateral de acento, título y hasta 5 viñetas de máx. 10 palabras."""
    _check_bullets(bullets)
    _rect(slide, left, top, width, height, GS1_BG_LIGHT, line=GS1_BORDER)
    _rect(slide, left, top, 0.08, height, accent_color)
    _textbox(slide, left + 0.2, top + 0.1, width - 0.3, 0.45, title, 16, accent_color, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.6), Inches(width - 0.3),
                                  Inches(height - 0.7))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = "• " + b
        _style(r, 14, GS1_TEXT)
    return tb


def add_gs1_table(slide, left, top, width, height, headers, rows, col_widths=None):
    """Tabla: cabecera GS1_BLUE/blanco, filas alternas blanco y GS1_BG_LIGHT, celdas centradas vertical."""
    n_rows, n_cols = len(rows) + 1, len(headers)
    gt = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top), Inches(width),
                                Inches(height))
    table = gt.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    else:
        for i in range(n_cols):
            table.columns[i].width = Inches(width / n_cols)
    row_h = Inches(height / n_rows)
    for r in table.rows:
        r.height = row_h

    def fill_cell(cell, text, color, bg, bold):
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.1)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ""
        run = p.add_run()
        run.text = str(text)
        _style(run, 12, color, bold)

    for c, h in enumerate(headers):
        fill_cell(table.cell(0, c), h, WHITE, GS1_BLUE, True)
    for r, row in enumerate(rows, start=1):
        bg = WHITE if r % 2 == 1 else GS1_BG_LIGHT
        for c, val in enumerate(row):
            fill_cell(table.cell(r, c), val, GS1_TEXT, bg, False)
    return gt
